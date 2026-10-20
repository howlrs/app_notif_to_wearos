"""クライアント向けの説明資料（pptx、1 スライド 1 要点）を作る。

出力: docs/presentation/sitealert_poc_report.pptx
画像: docs/screenshots/watch/*.png（make_slides.py が作る時計画像）
内容の出典: docs/05 §2・§3（検証項目と結果）、docs/06 §2（決定事項）、docs/04（配信）
実行: python3 make_deck.py（python-pptx 1.0 が必要）
"""
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
WATCH = HERE.parent / "screenshots" / "watch"
OUT = HERE / "sitealert_poc_report.pptx"

FONT_LATIN = "BIZ UDPGothic"
FONT_EA = "BIZ UDPゴシック"

INK = RGBColor(0x1A, 0x1E, 0x23)
INK2 = RGBColor(0x4B, 0x55, 0x60)
INK3 = RGBColor(0x69, 0x72, 0x7D)
LINE = RGBColor(0xD3, 0xD8, 0xDE)
PANEL = RGBColor(0xF1, 0xF3, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ALERT = RGBColor(0xB3, 0x26, 0x1E)  # アプリの警告色に合わせる
ALERT_SOFT = RGBColor(0xE9, 0xB0, 0xAB)
DONE = RGBColor(0x1E, 0x7A, 0x34)
CAUTION = RGBColor(0x85, 0x56, 0x00)

SLIDE_W, SLIDE_H = 13.333, 7.5
LEFT = 0.75
WIDTH = SLIDE_W - LEFT * 2
TITLE = "工事現場向け スマートウォッチ通知システム　試作の動作確認報告"
# 資料本文（docs/）は和文と英数字の間に半角スペースを入れるが、スライドの大きな文字では間延びするため詰める
CJK_GAP = re.compile(r"(?<=[\u3000-\u30ff\u4e00-\u9fff\uff00-\uffef]) (?=[A-Za-z0-9])|(?<=[A-Za-z0-9%]) (?=[\u3000-\u30ff\u4e00-\u9fff\uff00-\uffef])")
SCREEN_NOTE = "画面: Wear OS 5（API 34）エミュレーターで試作アプリを動かして撮影（2026-10-19）"


# ---- 部品 ----

def set_font(run, size, bold=False, color=INK):
    font = run.font
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "ja-JP")
    for tag in ("a:latin", "a:ea", "a:cs"):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    for tag, face in (("a:latin", FONT_LATIN), ("a:ea", FONT_EA), ("a:cs", FONT_LATIN)):
        rpr.append(rpr.makeelement(qn(tag), {"typeface": face}))


def fill_text(tf, paras, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP, spacing=1.3, gap=0):
    """paras の各要素は文字列、(文字列, 書式) のリスト、または dict（runs と段落の書式）。"""
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        if not isinstance(para, dict):
            para = {"runs": para}
        runs = para["runs"]
        if isinstance(runs, str):
            runs = [(runs, {})]
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", align)
        p.line_spacing = para.get("spacing", spacing)
        if i > 0:
            p.space_before = Pt(para.get("before", gap))
        for text, style in runs:
            for j, line in enumerate(text.split("\n")):
                if j:
                    p.add_line_break()
                r = p.add_run()
                r.text = CJK_GAP.sub("", line)
                set_font(r, style.get("size", para.get("size", size)),
                         style.get("bold", para.get("bold", bold)),
                         style.get("color", para.get("color", color)))


def add_text(slide, x, y, w, h, paras, **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_text(box.text_frame, paras, **kw)
    return box


def plain(shape):
    """テーマ由来の影・線・色を外し、明示した書式だけにする。"""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    return shape


def add_box(slide, x, y, w, h, fill=PANEL, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            radius=0.12, line_w=1.0):
    sp = plain(slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = radius
    return sp


def add_line(slide, x1, y1, x2, y2, color=INK3, width=1.25, arrow=False, dash=False):
    c = plain(slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def add_watch(slide, name, x, y, w, alt):
    pic = slide.shapes.add_picture(str(WATCH / f"{name}.png"), Inches(x), Inches(y), width=Inches(w))
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def caption(slide, x, y, w, title, body, num=None):
    runs = [(title, {"bold": True, "size": 17})]
    if num is not None:
        runs.insert(0, (f"{num}  ", {"bold": True, "size": 17, "color": ALERT}))
    add_text(slide, x, y, w, 1.3, [
        {"runs": runs, "align": PP_ALIGN.CENTER},
        {"runs": body, "size": 13.5, "color": INK2, "align": PP_ALIGN.CENTER, "before": 6},
    ], spacing=1.3)


def bullets(slide, x, y, w, h, items, size=20, gap=16, label=None):
    paras = []
    if label:
        paras.append({"runs": label, "size": 15, "bold": True, "color": INK2})
    for k, item in enumerate(items):
        paras.append({"runs": item, "before": 10 if (label and k == 0) else gap})
    return add_text(slide, x, y, w, h, paras, size=size, spacing=1.35)


def cell_border(cell, bottom=None, width=9525):
    tcpr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB", "a:noFill", "a:solidFill"):
        for el in tcpr.findall(qn(tag)):
            tcpr.remove(el)
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tcpr.makeelement(qn(tag), {"w": str(width if tag == "a:lnB" and bottom else 0)})
        if tag == "a:lnB" and bottom is not None:
            fill = ln.makeelement(qn("a:solidFill"), {})
            fill.append(fill.makeelement(qn("a:srgbClr"), {"val": str(bottom)}))
            ln.append(fill)
        else:
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        tcpr.append(ln)
    tcpr.append(tcpr.makeelement(qn("a:noFill"), {}))


def add_table(slide, x, y, widths, header, rows, row_h=0.42, size=14, styles=None):
    """表の行は (セルの runs, ...)。styles は列ごとの既定書式。"""
    n_rows, n_cols = len(rows) + 1, len(widths)
    gf = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y),
                                Inches(sum(widths)), Inches(row_h * n_rows))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    style_id = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
    if style_id is not None:
        style_id.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    for c, w in enumerate(widths):
        tbl.columns[c].width = Inches(w)
    styles = styles or [{}] * n_cols
    for r in range(n_rows):
        tbl.rows[r].height = Inches(0.36 if r == 0 else row_h)
        for c in range(n_cols):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0 if c == 0 else 0.12)
            cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if r == 0:
                fill_text(cell.text_frame, [header[c]], size=12.5, bold=True, color=INK2,
                          anchor=MSO_ANCHOR.MIDDLE)
                cell_border(cell, bottom=INK3, width=12700)
            else:
                value = rows[r - 1][c]
                fill_text(cell.text_frame, [value if isinstance(value, list) else [(value, styles[c])]],
                          size=size, anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
                cell_border(cell, bottom=LINE)
    return gf


def new_slide(prs, headline, notes, footer=TITLE):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    title = slide.shapes.title
    title.left, title.top = Inches(LEFT), Inches(0.55)
    title.width, title.height = Inches(WIDTH), Inches(1.3)
    fill_text(title.text_frame, [[(headline, {})]], size=30, bold=True, spacing=1.2)
    page = len(prs.slides)
    add_text(slide, LEFT, 6.98, WIDTH - 1.0, 0.3, [footer], size=10, color=INK3)
    add_text(slide, SLIDE_W - LEFT - 1.0, 6.98, 1.0, 0.3, [str(page)], size=10, color=INK3,
             align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = CJK_GAP.sub("", notes)
    return slide


def use_font_in_theme(prs):
    """後から追記する文字やノートも BIZ UDPゴシックになるよう、テーマのフォントを差し替える。"""
    for master in (prs.slide_master, prs.notes_master):
        theme = master.part.part_related_by(RT.THEME)
        is_xml = hasattr(theme, "_element")  # ノート用のテーマは XmlPart、スライド用はバイト列
        root = theme._element if is_xml else etree.fromstring(theme.blob)
        for tag in ("a:majorFont", "a:minorFont"):
            for font in root.iter(qn(tag)):
                font.find(qn("a:latin")).set("typeface", FONT_LATIN)
                font.find(qn("a:ea")).set("typeface", FONT_EA)
        if not is_xml:
            theme._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ---- スライド ----

def s_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = slide.shapes.title
    title.left, title.top, title.width, title.height = Inches(LEFT), Inches(2.0), Inches(7.4), Inches(1.9)
    fill_text(title.text_frame, [[("工事現場向け\nスマートウォッチ通知システム", {})]],
              size=40, bold=True, spacing=1.2, anchor=MSO_ANCHOR.BOTTOM)
    add_text(slide, LEFT, 4.2, 7.0, 1.0, [
        "試作アプリの動作確認報告",
        {"runs": "Wear OS 5 エミュレーターでの検証結果と、次に決めること", "size": 16,
         "color": INK2, "before": 10},
    ], size=22, color=INK)
    add_text(slide, LEFT, 6.55, 6.0, 0.35, ["2026年10月19日"], size=13, color=INK3)
    add_watch(slide, "08_alert_screen", 8.05, 1.2, 4.9, "警告画面。赤い背景に「警告: 閾値超過」、送信時刻、現場 101 / 端末 05、「確認」ボタン")
    slide.notes_slide.notes_text_frame.text = CJK_GAP.sub("", (
        "既存の DB 監視システムが閾値超過を検知した時に、作業員のスマートウォッチを鳴らす仕組みの試作報告。"
        "画面はすべて Wear OS 5 のエミュレーター（Google 公式の仮想ウォッチ）で、実機に入れるものと同じアプリを動かして撮影した。"
        "通知は Firebase 経由で実際に送っている。"))


def s_purpose(prs):
    s = new_slide(prs, "監視システムが異常を検知したら、\n作業員の手首へ直接知らせる", (
        "既存の監視システムが閾値超過を検知した時、その現場に登録された時計をすべて鳴らす。"
        "騒音の多い現場でも気付けるよう、音だけでなく振動でも知らせ、確認されるまで繰り返す。"
        "時計はスマホと連動させず、現場の Wi-Fi だけで動く。"), footer=SCREEN_NOTE)
    items = [
        ("知らせる相手", "異常を検知した現場で、時計を着けている作業員全員"),
        ("知らせ方", "振動とアラーム音。「確認」が押されるまで繰り返す（最大 60 秒）"),
        ("作業員がすること", "手首を上げて内容を見て、「確認」を押す"),
    ]
    paras = []
    for k, (label, text) in enumerate(items):
        paras.append({"runs": label, "size": 14, "bold": True, "color": INK2, "before": 0 if k == 0 else 26})
        paras.append({"runs": text, "size": 21, "before": 4})
    add_text(s, LEFT, 2.3, 6.7, 4.2, paras, spacing=1.35)
    add_watch(s, "06_alert_received", 8.2, 1.75, 4.4, "通知の表示。「現場アラート」「警告: 閾値超過」と警告の本文")


def s_architecture(prs):
    s = new_slide(prs, "スマホも SIM も使わず、\n現場の Wi-Fi だけで通知が届く", (
        "送信は Google の通知配信サービス（Firebase Cloud Messaging）を使う。時計は現場の Wi-Fi からインターネットに出られれば受信できる。"
        "時計は起動時と番号の変更時に、時計固有の識別子・現場番号・端末番号と、通知の宛先 ID を登録 API に送る。"
        "送信処理は既存の監視システムに組み込む想定で、参考実装を用意している。"))
    add_text(s, LEFT, 1.95, WIDTH, 0.5,
             ["時計は現場の Wi-Fi につなぐだけで動く。スマホとのペア設定や SIM の契約はいらない。"],
             size=17, color=INK2)
    nodes = [
        ("既存の\n監視システム", "閾値超過を検知する"),
        ("通知の\n送信処理", "現場の時計すべてへ送る"),
        ("Firebase Cloud\nMessaging", "Google の配信サービス"),
        ("現場の\nWi-Fi", "インターネットへの出口"),
    ]
    box_w, box_h, gap, top = 1.95, 1.05, 0.5, 3.1
    centers = []
    for i, (name, note) in enumerate(nodes):
        x = LEFT + i * (box_w + gap)
        box = add_box(s, x, top, box_w, box_h, fill=PANEL)
        fill_text(box.text_frame, [[(name, {})]], size=16, bold=True, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE, spacing=1.15)
        add_text(s, x - 0.15, top + box_h + 0.3, box_w + 0.3, 0.4, [note], size=13, color=INK2,
                 align=PP_ALIGN.CENTER)
        add_line(s, x + box_w + 0.06, top + box_h / 2, x + box_w + gap - 0.06, top + box_h / 2,
                 color=INK2, width=1.5, arrow=True)
        centers.append(x + box_w / 2)
    wx = LEFT + 4 * (box_w + gap)
    add_watch(s, "05_registered_standby", wx + 0.1, top + box_h / 2 - 0.85, 1.7,
              "登録済みの時計。現場 101 / 端末 05（待機中）")
    add_text(s, wx - 0.2, top + box_h + 0.3, 2.4, 0.4, ["振動と音で知らせる"], size=13, color=INK2,
             align=PP_ALIGN.CENTER)
    centers.append(wx + 0.95)
    # 登録の流れ（時計 → 送信処理）
    y0, y1 = top + box_h + 0.75, 5.75
    add_line(s, centers[4], y0, centers[4], y1, color=INK3, dash=True)
    add_line(s, centers[4], y1, centers[1], y1, color=INK3, dash=True)
    add_line(s, centers[1], y1, centers[1], y0, color=INK3, dash=True, arrow=True)
    add_text(s, centers[1] + 0.3, y1 + 0.12, centers[4] - centers[1] - 0.6, 0.4,
             ["端末の登録（初回と番号の変更時）: 現場番号・端末番号と通知の宛先 ID"],
             size=13, color=INK2, align=PP_ALIGN.CENTER)


def s_setup(prs):
    s = new_slide(prs, "初期設定は、現場番号と端末番号を入れるだけ", (
        "キッティング（配布前の設定）で 1 台ずつ行う。入力は番号 2 つだけで、文字入力はない。"
        "1 の通知許可の画面が英語なのは、エミュレーターの OS 部品に日本語訳が入っていないため。実機では日本語で出る想定。"
        "番号の範囲は現場番号 1〜999999、端末番号 1〜999（暫定）。"), footer=SCREEN_NOTE)
    shots = [
        ("01_permission", "通知を許可", "初回の起動時だけ\n（英語表示はエミュレーターのため）", "通知の許可を求める画面"),
        ("02_register_empty", "番号の入力欄", "現場番号と端末番号の\n2 つだけ", "端末登録画面。現場番号と端末番号の入力欄"),
        ("03_keypad", "数字キーで入力", "数字だけのキーが出る", "数字だけのテンキー"),
        ("04_register_filled", "「登録 / 更新」", "番号がサーバーに\n登録される", "現場 101、端末 5 を入力し「登録 / 更新」ボタンを表示"),
    ]
    w = 2.55
    gap = (WIDTH - w * 4) / 3
    for i, (name, title, body, alt) in enumerate(shots):
        x = LEFT + i * (w + gap)
        add_watch(s, name, x, 2.05, w, alt)
        caption(s, x - 0.15, 4.75, w + 0.3, title, body, num=i + 1)


def s_standby(prs):
    s = new_slide(prs, "登録後は、現場番号と端末番号が\n大きく表示される", (
        "待機画面には常に現場番号と端末番号を出し、取り違えを防ぐ。"
        "状態は、待機中（登録済み）、送信中（登録を送っている）、未同期（送れていない。自動で再送する）の 3 つ。"
        "番号の変更は「設定変更」の長押しだけで、触れただけでは入力画面に戻らない。"), footer=SCREEN_NOTE)
    bullets(s, LEFT, 2.3, 6.8, 4.2, [
        "どの現場の何番の端末か、時計を見ればすぐ分かる。",
        "状態は「待機中」「送信中」「未同期」の 3 つで表示する。",
        "番号の変更は「設定変更」を長押しした時だけ。触れただけでは変わらない。",
    ], size=21, gap=22)
    add_watch(s, "05_registered_standby", 8.2, 1.75, 4.4, "待機画面。現場: 101、端末: 05、（待機中）、設定変更ボタン")


def s_duplicate(prs):
    s = new_slide(prs, "番号を打ち間違えても、\n通知が届かなくなる時計は出ない", (
        "朝会の後に時計を配り、作業員が貸し出し台帳を見ながら現場番号と端末番号を入力する運用を想定している。"
        "サーバーは時計ごとの固有の識別子で登録を管理し、作業員が入力した端末番号は台帳との照合に使う。"
        "そのため、同じ日に別の時計と同じ番号を入力しても、先に登録した時計の登録は消えず、両方に通知が届く。"
        "重なった時計には「同じ番号の時計あり　台帳を確認」と表示し、管理側の一覧でも重なりが分かる。"
        "前日までの登録と番号が重なるのは貸し出しの入れ替えとして扱う（業務日は 4 時区切り、設定で変更可）。"
        "返却・紛失した時計は管理側で登録から外す。"), footer=SCREEN_NOTE)
    bullets(s, LEFT, 2.3, 6.8, 4.2, [
        "時計ごとの固有の識別子で登録するので、同じ番号を入れても先の時計の登録は消えない。",
        "通知は現場の時計すべてに届く。番号は貸し出し台帳との照合に使う。",
        "同じ日に番号が重なった時計には「同じ番号の時計あり　台帳を確認」と出る。",
    ], size=20, gap=22)
    add_watch(s, "11_duplicate_warning", 8.2, 1.75, 4.4,
              "待機画面。現場: 101、端末: 05、（待機中）、同じ番号の時計あり 台帳を確認")


def s_speed(prs):
    s = new_slide(prs, "画面が消えた省電力状態でも、\n送信から約 1 秒で鳴り始めた", (
        "画面を消し、省電力状態（Doze）を強制した上で、Firebase 経由で実際に送った。計測した回では約 1 秒で振動とアラーム音が始まった。"
        "エミュレーターでの値で、現場の Wi-Fi と実機での到達時間は別に測る必要がある（実機確認の項目に含めた）。"))
    x0, x1, y = 1.6, 11.7, 3.85
    add_line(s, x0, y, x1, y, color=INK2, width=2)
    for x, fill in ((x0, INK), (x1, ALERT)):
        add_box(s, x - 0.17, y - 0.17, 0.34, 0.34, fill=fill, shape=MSO_SHAPE.OVAL)
    add_text(s, x0, 2.55, x1 - x0, 1.0, ["約 1 秒"], size=54, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.BOTTOM, spacing=1.0)
    for x, title, body, align in (
        (x0 - 0.85, "送信", "送信処理から Firebase 経由で送る", PP_ALIGN.LEFT),
        (x1 - 2.65, "鳴動開始", "振動とアラーム音が始まる", PP_ALIGN.RIGHT),
    ):
        add_text(s, x, y + 0.4, 3.5, 1.0, [
            {"runs": title, "size": 18, "bold": True},
            {"runs": body, "size": 14, "color": INK2, "before": 4},
        ], align=align)
    add_line(s, LEFT, 5.55, LEFT + WIDTH, 5.55, color=LINE, width=1)
    add_text(s, LEFT, 5.75, WIDTH, 0.9, [
        [("測定条件　", {"bold": True, "color": INK2}),
         ("Wear OS 5 エミュレーター、画面消灯、省電力状態（Doze）を強制、Firebase 経由の実送信。計測した回の値。", {})],
        {"runs": "現場の Wi-Fi と実機での到達時間は、実機確認で測る。", "before": 6},
    ], size=14, color=INK2)


def s_repeat(prs):
    s = new_slide(prs, "「確認」されるまで、\n振動とアラーム音を最大 60 秒繰り返す", (
        "振動は「1 秒振動＋0.5 秒休止」の繰り返し、音はアラームの音として鳴らし続ける。"
        "騒音の中で一度きりの短い振動を逃さないよう、確認まで繰り返す方式に決めた（当初案は約 7 秒で 1 回）。"
        "60 秒で止まった後も通知は残り、手首を上げれば内容を確認できる。おやすみモード時の音の扱いは実機で確認する。"))
    x0, x1 = 3.35, 12.35
    per_sec = (x1 - x0) / 60
    rows = [
        ("確認されない場合", 3.0, 60, "60 秒で自動停止。通知は残る"),
        ("「確認」を押した場合", 4.35, 10.5, "押した時点で止まり、通知も消える"),
    ]
    add_text(s, x0, 2.38, 6.0, 0.35, ["1 秒振動＋0.5 秒休止を繰り返す"], size=13, color=INK2)
    for label, y, stop, note in rows:
        add_text(s, LEFT, y + 0.05, 2.45, 0.6, [label], size=16, bold=True)
        t = 0.0
        while t + 1.0 <= stop + 1e-6:
            add_box(s, x0 + t * per_sec, y, per_sec * 1.0, 0.42, fill=ALERT, shape=MSO_SHAPE.RECTANGLE)
            t += 1.5
        add_box(s, x0, y + 0.5, stop * per_sec, 0.1, fill=ALERT_SOFT, shape=MSO_SHAPE.RECTANGLE)
        xs = x0 + stop * per_sec
        add_line(s, xs, y - 0.12, xs, y + 0.72, color=INK, width=1.75)
        if stop < 60:
            add_text(s, xs + 0.15, y + 0.08, 5.0, 0.4, [note], size=14, color=INK)
        else:
            add_text(s, xs - 5.0, y + 0.72, 5.0, 0.35, [note], size=14, color=INK, align=PP_ALIGN.RIGHT)
    axis_y = 5.45
    add_line(s, x0, axis_y, x1, axis_y, color=LINE, width=1)
    for sec in (0, 15, 30, 45, 60):
        x = x0 + sec * per_sec
        add_line(s, x, axis_y, x, axis_y + 0.08, color=INK3, width=1)
        add_text(s, x - 0.5, axis_y + 0.12, 1.0, 0.3, [f"{sec} 秒"], size=12, color=INK3, align=PP_ALIGN.CENTER)
    for k, (fill, h, text) in enumerate(((ALERT, 0.2, "振動"), (ALERT_SOFT, 0.08, "アラーム音"))):
        y = 6.2 + k * 0.32
        add_box(s, x0, y + (0.2 - h) / 2, 0.3, h, fill=fill, shape=MSO_SHAPE.RECTANGLE)
        add_text(s, x0 + 0.42, y - 0.03, 2.0, 0.3, [text], size=12, color=INK2)


def s_wrist(prs):
    s = new_slide(prs, "手首を上げると、\n警告の内容と「確認」ボタンが表示される", (
        "鳴動中に手首を上げると、通知として警告の題名と本文が出る。"
        "下へ送ると「確認」「スマートウォッチで開く」「クリア」が並ぶ。「確認」で鳴動が止まる。"
        "「開く」を押すとアプリの警告画面になり、内容・送信時刻・現場と端末の番号、大きな「確認」ボタンが出る。"),
        footer=SCREEN_NOTE)
    shots = [
        ("06_alert_received", "通知の表示", "警告の題名と本文が出る", "通知の表示。警告: 閾値超過と本文"),
        ("07_alert_actions", "下へ送ると操作", "「確認」で止める\n「開く」で警告画面へ", "通知の操作。確認、スマートウォッチで開く、クリア"),
        ("08_alert_screen", "警告画面", "内容・送信時刻・番号と\n大きな「確認」ボタン", "警告画面。送信時刻と現場・端末の番号、確認ボタン"),
    ]
    w = 3.2
    gap = (WIDTH - w * 3) / 2
    for i, (name, title, body, alt) in enumerate(shots):
        x = LEFT + i * (w + gap)
        add_watch(s, name, x, 1.95, w, alt)
        caption(s, x, 5.2, w, title, body, num=i + 1)
        if i < 2:
            mid = 1.95 + w * 0.5
            add_line(s, x + w + 0.12, mid, x + w + gap - 0.12, mid, color=INK3, width=1.5, arrow=True)


def s_ack(prs):
    s = new_slide(prs, "「確認」を押すと、振動と音はすぐに止まる", (
        "警告画面または通知の「確認」を押すと、振動・アラーム音が止まり、通知も消える。"
        "押してから約 1 秒後に確かめた時点で、どれも止まっていた。時計は通常の文字盤に戻る。"),
        footer=SCREEN_NOTE)
    w = 3.7
    xa, xb = 1.7, SLIDE_W - 1.7 - w
    add_watch(s, "08_alert_screen", xa, 1.95, w, "警告画面と「確認」ボタン")
    add_watch(s, "09_after_ack", xb, 1.95, w, "確認後に戻った時計の文字盤")
    mid = 1.95 + w * 0.5
    add_line(s, xa + w + 0.2, mid, xb - 0.2, mid, color=INK2, width=2, arrow=True)
    add_text(s, xa + w + 0.2, mid - 0.55, xb - xa - w - 0.4, 0.4, ["「確認」を押す"], size=17, bold=True,
             align=PP_ALIGN.CENTER)
    caption(s, xa, 5.35, w, "鳴動中の警告画面", "振動とアラーム音が続いている")
    caption(s, xb, 5.35, w, "通常の文字盤に戻る", "振動・音・通知がすべて止まる\n（約 1 秒後に確認）")


def s_test(prs):
    s = new_slide(prs, "配る前に、テスト通知で\n1 台ずつ鳴動を確かめられる", (
        "キッティングの最後に、管理側の送信ツールから現場番号と端末番号を指定してテスト通知を送る。"
        "題名に【テスト】が付くので、本番の警告と見分けられる。鳴り方と「確認」の操作は本番と同じ。"),
        footer=SCREEN_NOTE)
    bullets(s, LEFT, 2.3, 6.8, 4.2, [
        "管理側の送信ツールで、現場番号と端末番号を指定して送る。",
        "題名の先頭に【テスト】が付き、本番の警告と見分けられる。",
        "鳴り方と「確認」の操作は本番と同じ。キッティングの手順に組み込んだ。",
    ], size=21, gap=22)
    add_watch(s, "10_test_notification", 8.2, 1.75, 4.4, "テスト通知。【テスト】鳴動確認")


def s_results(prs):
    s = new_slide(prs, "エミュレーターでは、\n登録から警告の停止まで一通り動いた", (
        "Wear OS 5 エミュレーターで、Firebase 経由の実送信を使って確認した結果。"
        "警告の自動表示だけは想定どおりに動かず、手首を上げて表示する方式で対応した（次のページ）。"
        "おやすみモード時の音、電池持ち、騒音下での気付き、Google Play からの配信は未確認。"
        "アプリの単体テスト 19 件、送信処理のテスト 34 件はすべて成功。"))
    ok = ("確認済み", {"bold": True, "color": DONE})
    rows = [
        ("端末の登録", "番号と宛先がサーバーに保存され、画面が「待機中」になった", ok),
        ("画面消灯・省電力中の受信", "受信して鳴り始めた（計測した回は送信から約 1 秒）", ok),
        ("振動とアラーム音", "1 秒振動＋0.5 秒休止の繰り返しと、アラーム音の再生", ok),
        ("自動停止", "確認されないと 60 秒で止まり、通知は残る", ok),
        ("「確認」での停止", "振動と音が止まり、通知も消えた", ok),
        ("切断中に送った警告", "接続が戻った後、保持期間（5 分）内に届いた", ok),
        ("アプリ更新後の受信", "上書き更新の後も、登録し直さずに受信した", ok),
        ("番号の入力", "数字だけのキーが出た", ok),
        ("警告の自動表示", "自動では全画面にならない。手首を上げると表示された", ("代替で対応", {"bold": True, "color": CAUTION})),
    ]
    add_table(s, LEFT, 1.95, [3.2, 6.9, 1.73], ["確認項目", "結果", "判定"],
              [(a, b, [c]) for a, b, c in rows], row_h=0.43, size=14,
              styles=[{"bold": True}, {}, {}])
    add_text(s, LEFT, 6.35, WIDTH, 0.4,
             ["未確認: おやすみモード時の音、電池持ち、騒音下での気付き、Google Play からの配信（実機で確認する）"],
             size=13, color=INK2)


def s_fullscreen(prs):
    s = new_slide(prs, "Wear OS 5 では、\n警告が自動で全画面には表示されない", (
        "当初は、受信と同時に画面を点けて警告画面を全面に出す想定だった（全画面通知）。"
        "Wear OS 5 のエミュレーターでは、OS 側でこの仕組みが無効になっていて起動しなかった。"
        "そこで、振動と音で気付かせ、作業員が手首を上げて内容を見る運用を前提にした。"
        "実機（機種ごと）でも同じ動きかを確かめる。"))
    add_text(s, LEFT, 1.95, WIDTH, 0.8, [
        "当初は、受信と同時に画面を点けて警告画面を全面に出す想定だった。"
        "Wear OS 5 のエミュレーターでは、この仕組み（全画面通知）が OS 側で働かなかった。"
    ], size=17, color=INK2, spacing=1.4)
    add_text(s, LEFT, 3.05, WIDTH, 0.4, ["対応：気付かせるのは振動と音、内容は手首を上げて見る"],
             size=18, bold=True)
    steps = [
        ("時計", "振動とアラーム音で\n知らせる（自動）", ALERT, None),
        ("作業員", "気付いて\n手首を上げる", WHITE, INK),
        ("時計", "警告の内容と\n「確認」ボタンを出す", INK, None),
    ]
    y = 4.15
    xs = [LEFT + 1.4, SLIDE_W / 2, SLIDE_W - LEFT - 1.4]
    add_line(s, xs[0], y, xs[2], y, color=LINE, width=2)
    for x, (who, what, fill, line) in zip(xs, steps):
        add_box(s, x - 0.2, y - 0.2, 0.4, 0.4, fill=fill, line=line, shape=MSO_SHAPE.OVAL, line_w=2)
        add_text(s, x - 1.6, y + 0.38, 3.2, 1.4, [
            {"runs": who, "size": 13, "bold": True, "color": INK2},
            {"runs": what, "size": 18, "bold": True, "before": 4},
        ], align=PP_ALIGN.CENTER, spacing=1.3)
    add_text(s, LEFT, 6.3, WIDTH, 0.4, ["実機（機種ごと）でも同じ動きかを確かめる。"], size=14, color=INK2)


def s_network(prs):
    s = new_slide(prs, "届くまでの時間は、\n現場の Wi-Fi と実機で測る必要がある", (
        "開発環境では、長時間操作しないでいると通知を受ける接続が切れたままになり、つなぎ直すまで届かないことがあった。"
        "開発用の回線（VPN 経由）に特有の現象の可能性があるが、現場の Wi-Fi で同じことが起きないかを測る。"
        "届かなかった警告は 5 分間だけ保持し、それより古い警告は後から鳴らさない設定にしている。"))
    col_w = (WIDTH - 0.8) / 2
    bullets(s, LEFT, 2.15, col_w, 4.3, [
        "長い時間なにもしないでいると、通知を受ける接続が切れたままになり、つなぎ直すまで届かないことがあった。",
        "開発用の回線（VPN 経由）に特有の現象の可能性がある。",
    ], size=18, gap=14, label="開発環境で起きたこと")
    add_line(s, LEFT + col_w + 0.4, 2.2, LEFT + col_w + 0.4, 6.3, color=LINE, width=1)
    bullets(s, LEFT + col_w + 0.8, 2.15, col_w, 4.3, [
        "1 時間以上待機した後でも、10 秒以内に届くか",
        "Wi-Fi の省電力機能で、接続が切れたままにならないか",
        "切れていた間の警告が、つながった後に届くか（5 分を過ぎた古い警告は鳴らさない設定）",
    ], size=18, gap=14, label="現場で測ること")


def s_device(prs):
    s = new_slide(prs, "次は実機で、4 つの点を確かめる", (
        "エミュレーターでは確かめられない項目。合格の目安は暫定で、現場試行の前に見直す。"
        "検証用に実機 2 台を調達する想定。"))
    rows = [
        ("画面の表示", "手首を上げた時に、警告と「確認」ボタンが出る。自動で点灯するかも記録する"),
        ("届くまでの時間", "画面を消して 1 時間以上置いた後でも、送信から 10 秒以内に通知が出る"),
        ("騒音下での気付き", "重機の近くなど実際の騒音の中で、作業員が振動と音に気付ける"),
        ("電池持ち", "Wi-Fi につないだまま 1 勤務（8〜10 時間）着けて、終業時に 20% 以上残る"),
    ]
    add_table(s, LEFT, 2.05, [3.0, WIDTH - 3.0], ["確認すること", "合格の目安（暫定）"], rows,
              row_h=0.85, size=17, styles=[{"bold": True}, {}])


def s_distribution(prs):
    s = new_slide(prs, "アプリの配布と更新は、\nGoogle Play の内部テストで行う", (
        "Google Play の「内部テスト」に登録したアカウントだけが、時計の Play ストアからインストールできる。更新も同じ経路で届く。"
        "テスターは最大 100 アカウント。少数の共通アカウントで時計にログインする運用にした。"
        "開始には Google Play Console の開発者登録（登録料 25 米ドル、本人確認）が要る。まだ登録していない。"))
    bullets(s, LEFT, 2.3, 7.0, 4.2, [
        "時計の Play ストアからインストールでき、更新も同じ経路で届く。",
        "テスターは最大 100 アカウント。少数の共通アカウントで時計にログインして使う。",
        [("開始には Google Play Console の開発者登録（登録料 25 米ドル、本人確認）が要る。", {}),
         ("まだ登録していない。", {"bold": True, "color": CAUTION})],
    ], size=20, gap=22)
    flow = ["アプリを\n内部テストに公開", "Google Play", "テスター用アカウントで\nログインした時計"]
    x, w, h = 8.85, 3.7, 0.95
    for i, text in enumerate(flow):
        y = 2.3 + i * (h + 0.5)
        box = add_box(s, x, y, w, h, fill=PANEL)
        fill_text(box.text_frame, [[(text, {})]], size=16, bold=True, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE, spacing=1.15)
        if i < 2:
            add_line(s, x + w / 2, y + h + 0.06, x + w / 2, y + h + 0.44, color=INK2, width=1.5, arrow=True)


def s_existing(prs):
    s = new_slide(prs, "既存の監視システムへ組み込むため、\n構成を教えてください", (
        "登録 API と送信処理は、既存の監視システムと同じ環境に置くのが自然。言語や DB に合わせて移植する。"
        "現場の Wi-Fi から登録 API へ届くこと、HTTPS で公開するドメインが要る。"
        "異常が続く間の再通知は、既存システムの警報ルールに合わせる。"))
    col_w = 6.4
    bullets(s, LEFT, 2.35, col_w, 4.0, [
        "使っているプログラム言語とデータベース製品",
        "サーバーの置き場所（現場の Wi-Fi から登録 API へ届くか）",
        "登録 API に使うドメイン名（HTTPS で公開）",
        "異常が続いた時に、警報を繰り返すルール",
    ], size=19, gap=14, label="教えていただきたいこと")
    px = LEFT + col_w + 0.6
    add_box(s, px, 2.0, SLIDE_W - LEFT - px, 3.9, fill=PANEL, radius=0.05)
    bullets(s, px + 0.4, 2.35, SLIDE_W - LEFT - px - 0.8, 3.4, [
        "登録 API と送信処理の\n参考実装（Python）",
        "テーブル定義\n（PostgreSQL・MySQL 用）",
        "自動テスト 34 件（すべて成功）",
    ], size=18, gap=14, label="こちらで用意済み")


def s_decisions(prs):
    s = new_slide(prs, "対象機種は、現場試行の前に決めたい", (
        "対象機種は、スピーカー付き、Wear OS 4 以降、2.4GHz の Wi-Fi に対応していることを条件にしたい。"
        "通知の宛先は現場単位に決まった（現場の時計すべてへ送る）。"
        "確認結果の表示と端末の生存監視は、試行の結果を見て決める。"))
    col_w = (WIDTH - 0.8) / 2
    bullets(s, LEFT, 2.15, col_w, 4.3, [
        [("対象機種　", {"bold": True}), ("スピーカー付き、Wear OS 4 以降、2.4GHz の Wi-Fi に対応", {})],
        [("決定済み　", {"bold": True, "color": DONE}), ("通知は現場単位で、現場の時計すべてへ送る", {})],
    ], size=19, gap=18, label="現場試行の前に決める")
    add_line(s, LEFT + col_w + 0.4, 2.2, LEFT + col_w + 0.4, 6.3, color=LINE, width=1)
    bullets(s, LEFT + col_w + 0.8, 2.15, col_w, 4.3, [
        [("確認の記録　", {"bold": True}), ("「確認」したことを監視画面に表示するか", {})],
        [("端末の見守り　", {"bold": True}), ("動いていない端末を定期的に見つける仕組みが要るか", {})],
    ], size=19, gap=18, label="現場試行の結果で決める")


def s_roadmap(prs):
    s = new_slide(prs, "実機と配信の準備を進め、\n1 現場での試行につなげる", (
        "1 と 2 は並行して進められる。4 は既存システムの構成が分かってから着手する。"
        "試行は 1 現場で 2 週間ほどを想定し、到達率・遅延・電池・作業員の声をまとめて本格展開を判断する。"))
    steps = [
        ("済", "試作", "アプリと送信処理\nエミュレーター確認"),
        ("1", "実機の調達", "機種を選び\n検証用に 2 台"),
        ("2", "配信の準備", "Play Console 登録\n内部テスト配信"),
        ("3", "実機での確認", "4 項目を確かめる"),
        ("4", "組み込み", "既存システムへ\n結合テスト"),
        ("5", "キッティング", "2〜3 台で\n手順を試す"),
        ("6", "現場試行", "1 現場で\n2 週間ほど"),
    ]
    y = 3.55
    x0, x1 = LEFT + 0.8, SLIDE_W - LEFT - 0.8
    step = (x1 - x0) / (len(steps) - 1)
    add_line(s, x0, y, x1, y, color=LINE, width=2)
    for i, (mark, title, detail) in enumerate(steps):
        x = x0 + i * step
        done = mark == "済"
        dot = add_box(s, x - 0.3, y - 0.3, 0.6, 0.6, fill=DONE if done else WHITE,
                      line=None if done else INK, shape=MSO_SHAPE.OVAL, line_w=1.75)
        fill_text(dot.text_frame, [[(mark, {})]], size=16, bold=True, color=WHITE if done else INK,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
        add_text(s, x - 0.85, y + 0.5, 1.7, 2.0, [
            {"runs": title, "size": 16, "bold": True, "color": DONE if done else INK},
            {"runs": detail, "size": 13, "color": INK2, "before": 6},
        ], align=PP_ALIGN.CENTER, spacing=1.35)


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    use_font_in_theme(prs)
    props = prs.core_properties
    props.title = "工事現場向け スマートウォッチ通知システム 試作の動作確認報告"
    props.author = ""
    props.last_modified_by = ""
    props.language = "ja-JP"
    for build in (s_cover, s_purpose, s_architecture, s_setup, s_standby, s_duplicate, s_speed, s_repeat,
                  s_wrist, s_ack, s_test, s_results, s_fullscreen, s_network, s_device,
                  s_distribution, s_existing, s_decisions, s_roadmap):
        build(prs)
    prs.save(OUT)
    print(f"{OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
